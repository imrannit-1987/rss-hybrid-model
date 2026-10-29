"""Workbook: month-by-month model vs actual for all sites. usage: sim_all_workbook.py sim_all_monthly.csv out.xlsx"""
import sys, pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
R=pd.read_csv(sys.argv[1],parse_dates=['month']).sort_values(['project','month']); out=sys.argv[2]
ALL=R.copy(); R=R[R.flag=='ok'].copy()
F='Arial'; H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E5F'); red=PatternFill('solid',fgColor='F8CBAD'); grn=PatternFill('solid',fgColor='C6E0B4')
def hdr(ws,cols):
    for j,c in enumerate(cols,1):
        x=ws.cell(row=1,column=j,value=c); x.font=H; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=42; ws.freeze_panes='A2'
def autow(ws):
    for col in ws.columns:
        L=max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200]); ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(9,L+2),34)
wb=Workbook()
# Overview
ws=wb.active; ws.title='Overview'
ws['A1']='Model vs actual, every site, every month with a month-end report'; ws['A1'].font=Font(name=F,bold=True,size=13)
n=len(R)+1
ws['A3']='Site-months'; ws['B3']=f"=COUNTA('Site-months'!A2:A{n})"
ws['A4']='Sites'; ws['B4']=R.project.nunique()
ws['A5']='Monthly L/kWh error, median of |error|'; ws['B5']=f"=MEDIAN('Site-months'!R2:R{n})"
ws['A6']='Monthly L/kWh error, 90th percentile of |error|'; ws['B6']=f"=PERCENTILE('Site-months'!R2:R{n},0.9)"
ws['A7']='Share of months within ±10%'; ws['B7']=f"=COUNTIF('Site-months'!R2:R{n},\"<=0.1\")/COUNTA('Site-months'!R2:R{n})"
ws['A8']='Bias (median error, + = model high)'; ws['B8']=f"=MEDIAN('Site-months'!Q2:Q{n})"
ws['A9']='Fleet fuel, actual (L)'; ws['B9']=f"=SUM('Site-months'!H2:H{n})"
ws['A10']='Fleet fuel, model (L)'; ws['B10']=f"=SUM('Site-months'!L2:L{n})"
ws['A11']='Fleet fuel error'; ws['B11']='=B10/B9-1'
for r in range(3,12): ws[f'A{r}'].font=Font(name=F); ws[f'B{r}'].font=Font(name=F,bold=True)
for r in (5,6,7,8,11): ws[f'B{r}'].number_format='0.0%'
for r in (9,10): ws[f'B{r}'].number_format='#,##0'
ws['A13']='How to read'; ws['A13'].font=Font(name=F,bold=True)
notes=['Each site-month was simulated from the site\'s hourly generator demand (unit counters, scaled to the month-end DG kWh so the kWh basis is the site meter).',
'The PMS start/stop settings come from the site\'s own 1-minute data for that quarter (pms_source = quarter), else the site median, else a default by site type.',
'Fuel per unit comes from the engine-family curves fitted with THAT SITE LEFT OUT, so every number is an out-of-sample prediction.',
'"Curves on actual running" applies the same curves to what the units really did; the gap between it and the model is the dispatch model\'s own error.',
'Actual = month-end Hybrid Performance Report fuel and DG kWh. Months with fewer than 20 days of hourly data, or where the report DG differs from the portal meter by more than 20% (site changes, partial billing months), are listed on the Excluded months tab and left out of the statistics.']
for i,t in enumerate(notes): ws[f'A{14+i}']=t; ws[f'A{14+i}'].font=Font(name=F,size=10)
ws.column_dimensions['A'].width=48; ws.column_dimensions['B'].width=16
# Site-months
wm=wb.create_sheet('Site-months')
cols=['Site','Month','Country','Type','Days of hourly data','Fleet (kVA)','Engine families','Actual fuel (L)','Actual DG kWh','Actual L/kWh','Solar share of load','Model fuel (L)','Model L/kWh','Curves on actual running (L)','PMS start','PMS stop','Model error','|Model error|','PMS setting source','Sim mean loading','Site-meter ratio (report DG / unit counters)']
hdr(wm,cols)
for i,x in enumerate(R.itertuples(),2):
    vals=[x.project,x.month.strftime('%Y-%m'),x.country,x.type,int(x.days_with_data),x.fleet,x.families,round(x.act_fuel),round(x.dg),None,round(x.solar_share,3) if pd.notna(x.solar_share) else None,round(x.model_fuel),None,round(x.curves_on_actual_fuel),round(x.pms_start,2),round(x.pms_stop,2),None,None,x.pms_source,round(x.sim_mean_loading,2) if pd.notna(x.sim_mean_loading) else None,round(x.meter_ratio,3)]
    for j,v in enumerate(vals,1): c=wm.cell(row=i,column=j,value=v); c.font=Font(name=F)
    wm.cell(row=i,column=10,value=f'=H{i}/I{i}'); wm.cell(row=i,column=13,value=f'=L{i}/I{i}'); wm.cell(row=i,column=17,value=f'=L{i}/H{i}-1'); wm.cell(row=i,column=18,value=f'=ABS(Q{i})')
    for j in (10,13): wm.cell(row=i,column=j).number_format='0.000'
    for j in (8,9,12,14): wm.cell(row=i,column=j).number_format='#,##0'
    for j in (11,17,18): wm.cell(row=i,column=j).number_format='0.0%'
    for j in (10,13,17,18): wm.cell(row=i,column=j).font=Font(name=F)
wm.conditional_formatting.add(f'R2:R{n}',CellIsRule(operator='greaterThan',formula=['0.1'],fill=red))
wm.conditional_formatting.add(f'R2:R{n}',CellIsRule(operator='lessThanOrEqual',formula=['0.05'],fill=grn))
wm.auto_filter.ref=f'A1:U{n}'; autow(wm)
# Excluded months
we=wb.create_sheet('Excluded months'); hdr(we,['Site','Month','Reason','Days of hourly data','Report DG kWh','Portal DG kWh','Actual L/kWh','Model L/kWh (unreliable)'])
for i,x in enumerate(ALL[ALL.flag!='ok'].itertuples(),2):
    for j,v in enumerate([x.project,x.month.strftime('%Y-%m'),x.flag,int(x.days_with_data),round(x.dg),round(x.dg_portal) if pd.notna(x.dg_portal) else None,round(x.act_LkWh,3),round(x.model_LkWh,3)],1): c=we.cell(row=i,column=j,value=v); c.font=Font(name=F)
autow(we)
# Site summary
S=R.groupby('project').agg(months=('month','size'),act_fuel=('act_fuel','sum'),dg=('dg','sum'),model_fuel=('model_fuel','sum'),curves_act=('curves_on_actual_fuel','sum'),first=('month','min'),last=('month','max')).reset_index()
S['med_abs']=R.groupby('project').model_err.apply(lambda e:e.abs().median()).values
wsx=wb.create_sheet('Site summary'); hdr(wsx,['Site','Months','From','To','Actual fuel (L)','Actual L/kWh','Model fuel (L)','Model L/kWh','Whole-period error','Typical monthly |error|','Dispatch-only error (model vs curves on actual running)'])
for i,x in enumerate(S.itertuples(),2):
    for j,v in enumerate([x.project,int(x.months),x.first.strftime('%Y-%m'),x.last.strftime('%Y-%m'),round(x.act_fuel),None,round(x.model_fuel),None,None,round(x.med_abs,4),None],1): c=wsx.cell(row=i,column=j,value=v); c.font=Font(name=F)
    wsx.cell(row=i,column=6,value=f"=E{i}/SUMIFS('Site-months'!$I$2:$I${n},'Site-months'!$A$2:$A${n},A{i})"); wsx.cell(row=i,column=8,value=f"=G{i}/SUMIFS('Site-months'!$I$2:$I${n},'Site-months'!$A$2:$A${n},A{i})")
    wsx.cell(row=i,column=9,value=f'=G{i}/E{i}-1'); wsx.cell(row=i,column=11,value=f"=G{i}/SUMIFS('Site-months'!$N$2:$N${n},'Site-months'!$A$2:$A${n},A{i})-1")
    for j in (6,8): wsx.cell(row=i,column=j).number_format='0.000'
    for j in (5,7): wsx.cell(row=i,column=j).number_format='#,##0'
    for j in (9,10,11): wsx.cell(row=i,column=j).number_format='0.0%'
    for j in (6,8,9,11): wsx.cell(row=i,column=j).font=Font(name=F)
m=len(S)+1; wsx.conditional_formatting.add(f'I2:I{m}',CellIsRule(operator='greaterThan',formula=['0.1'],fill=red)); wsx.conditional_formatting.add(f'I2:I{m}',CellIsRule(operator='lessThan',formula=['-0.1'],fill=red))
wsx.auto_filter.ref=f'A1:K{m}'; autow(wsx)
# Fleet by month
M=R.groupby('month').agg(sites=('project','nunique'),act_fuel=('act_fuel','sum'),dg=('dg','sum'),model_fuel=('model_fuel','sum')).reset_index()
wf=wb.create_sheet('Fleet by month'); hdr(wf,['Month','Sites','Actual fuel (L)','Actual DG kWh','Actual L/kWh','Model fuel (L)','Model L/kWh','Fleet error'])
for i,x in enumerate(M.itertuples(),2):
    for j,v in enumerate([x.month.strftime('%Y-%m'),int(x.sites),round(x.act_fuel),round(x.dg),None,round(x.model_fuel),None,None],1): c=wf.cell(row=i,column=j,value=v); c.font=Font(name=F)
    wf.cell(row=i,column=5,value=f'=C{i}/D{i}'); wf.cell(row=i,column=7,value=f'=F{i}/D{i}'); wf.cell(row=i,column=8,value=f'=F{i}/C{i}-1')
    for j in (5,7): wf.cell(row=i,column=j).number_format='0.000'
    for j in (3,4,6): wf.cell(row=i,column=j).number_format='#,##0'
    wf.cell(row=i,column=8).number_format='0.0%'
    for j in (5,7,8): wf.cell(row=i,column=j).font=Font(name=F)
autow(wf)
wb.save(out); print(out)
